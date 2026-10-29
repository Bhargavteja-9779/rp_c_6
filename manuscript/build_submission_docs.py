"""Title page and cover letter (python-docx). Author-specific fields are explicit placeholders;
nothing about authors, affiliations, funding or originality is asserted on the authors' behalf."""
import json
import os
import sys

from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from numbers_from_results import compute_numbers  # noqa: E402

TITLE = ("Seen or new? Duplicate records, conformal coverage on novel inputs, and a novelty-conditional "
         "remedy for tabular benchmarks")


def _doc():
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    return d


def title_page():
    d = _doc()
    d.add_heading("Title page", level=1)
    d.add_paragraph().add_run(TITLE).bold = True
    for label in ["Author names (in final order)", "Affiliations (department, institution, city, country)",
                  "Corresponding author and active e-mail address", "ORCID iDs (16 digits, if available)",
                  "Acknowledgments", "Funding statement", "Competing interests statement",
                  "Author contributions (CRediT)"]:
        p = d.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run("[AUTHOR TO COMPLETE]")
    d.add_paragraph("This title page is supplied separately so that the manuscript file can be reviewed anonymously; "
                    "the manuscript file itself contains no identifying information.")
    d.save(os.path.join(HERE, "title_page.docx"))


def cover_letter(n):
    d = _doc()
    d.add_paragraph("[DATE]")
    d.add_paragraph("Editor-in-Chief\nJournal of King Saud University – Computer and Information Sciences")
    d.add_paragraph("Dear Editor,")
    body = [
        f"We submit the manuscript “{TITLE}” for consideration as a research article in the Journal of King Saud "
        "University – Computer and Information Sciences.",
        "Conformal prediction is increasingly used to attach guaranteed uncertainty to machine-learning predictions, "
        "and its guarantee is usually checked on random row splits of benchmark data. The manuscript asks what this "
        "guarantee means when a dataset contains duplicated records, so that some test rows have exact copies in the "
        "training data. We derive an exact expression for coverage on genuinely new records, a label-free formula that "
        "predicts exposure from duplicate-group sizes, and a novelty-conditional (Mondrian) remedy that is valid for "
        "any mix of seen and new records.",
        f"A pre-registered evaluation on {n['n_total']} OpenML-CC18 and OpenML-CTR23 datasets, with five learners and "
        "four nonconformity scores, shows that the exposure formula is almost exact (mean absolute error "
        f"{n['exp_mae']}). It also shows that the primary hypothesis of a typical shortfall was not supported: median "
        f"gaps were {n['h1x_gap_median_min']} to {n['h1x_gap_median_max']} percentage points. Shortfalls are "
        f"concentrated in a predictable minority of datasets, reaching {n['h1_worst_gap']} points, and controlled "
        "duplicate-injection experiments reproduce the predicted dose–response. We report this negative primary "
        "result openly, together with the mechanism that explains it. All code, raw results and the analysis "
        "pipeline will be released.",
        "We believe the work suits the journal's readership because it addresses the foundations of reliable "
        "machine-learning evaluation and has direct consequences for practitioners who deploy uncertainty-aware "
        "models on tabular data.",
        "[AUTHOR TO CONFIRM: This manuscript is original, has not been published, and is not under consideration "
        "elsewhere. All authors have approved the submission. We declare no competing interests.] "
        "[AUTHOR TO COMPLETE: re-used or unpublished material, if any; suggested reviewers, if desired.]",
        "The manuscript was prepared with substantial assistance from an AI system, as described in its "
        "“Use of AI tools” statement.",
        "Sincerely,\n[CORRESPONDING AUTHOR NAME, AFFILIATION, E-MAIL]",
    ]
    for b in body:
        d.add_paragraph(b)
    d.save(os.path.join(HERE, "cover_letter.docx"))


def cover_letter_pr(n):
    d = _doc()
    d.add_paragraph("[DATE]")
    d.add_paragraph("The Editors-in-Chief\nPattern Recognition (Elsevier)")
    d.add_paragraph("Dear Editors,")
    body = [
        f"We submit the manuscript “{TITLE}” for consideration as a regular research article in Pattern Recognition.",
        "Conformal prediction has become a standard tool for distribution-free uncertainty quantification, and "
        "Pattern Recognition has published a substantial body of conformal methodology, including recent Mondrian "
        "approaches. Conformal guarantees are usually checked on random row splits of benchmark data. Our manuscript "
        "shows what such a guarantee means when records are duplicated: the coverage of genuinely new records is "
        "(c − πc_S)/(1 − π), where π is the share of test rows with a training copy and c_S is the coverage of those "
        "rows, falling to 1 − α/(1 − π) in the memorization limit. We give a label-free formula for π and show that a "
        "novelty-conditional Mondrian calibration is exactly valid for any mix of seen and new records.",
        f"The study follows a protocol fixed in version control before the full runs and covers {n['n_total']} "
        "OpenML-CC18/CTR23 datasets, five learners, four nonconformity scores and three miscoverage levels, plus "
        "controlled duplicate-injection experiments and an external validation on software-defect and intrusion-"
        f"detection data with documented duplication. The exposure formula is almost exact (mean absolute error "
        f"{n['exp_mae']}). We report openly that the pre-specified primary hypothesis was not supported for the typical "
        f"benchmark dataset (median gaps {n['h1x_gap_median_min']} to {n['h1x_gap_median_max']} percentage points). "
        f"Shortfalls are concentrated in a predictable minority of datasets, reaching {n['h1_worst_gap']} points in the "
        f"benchmark suites and {n['ext_scp_min']} points in the external data. The proposed calibration restores "
        "nominal coverage at negligible cost.",
        "We believe the paper will interest the journal's readers who develop or evaluate conformal and other "
        "uncertainty-quantification methods, because it identifies a common evaluation artefact, explains it exactly, "
        "and provides a practical remedy and audit. All code, raw results and the analysis pipeline will be released.",
        "[AUTHOR TO CONFIRM: The manuscript is original, has not been published, and is not under consideration "
        "elsewhere; all authors approve the submission; no competing interests.] [AUTHOR TO COMPLETE: suggested or "
        "excluded reviewers, if desired.]",
        "The manuscript was prepared with substantial assistance from an AI system, as declared in the manuscript.",
        "Sincerely,\n[CORRESPONDING AUTHOR NAME, AFFILIATION, E-MAIL]",
    ]
    for b in body:
        d.add_paragraph(b)
    d.save(os.path.join(HERE, "cover_letter_PatternRecognition.docx"))


TITLE_V2 = ("Duplicate Data Rows Make Prediction Uncertainty Look Better Than It Is: A Simple Check and Fix for "
            "New Cases")


def title_page_v2():
    d = _doc()
    d.add_heading("Title page", level=1)
    d.add_paragraph().add_run(TITLE_V2).bold = True
    for label in ["Author names (in final order)", "Affiliations (department, institution, city, country)",
                  "Corresponding author and active e-mail address", "ORCID iDs (16 digits, if available)",
                  "Acknowledgments", "Funding statement", "Competing interests statement",
                  "Author contributions (CRediT)"]:
        p = d.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run("[AUTHOR TO COMPLETE]")
    d.save(os.path.join(HERE, "Revised_Title_Page.docx"))


def cover_letter_v2(n):
    d = _doc()
    d.add_paragraph("[DATE]")
    d.add_paragraph("Editor-in-Chief\nJournal of King Saud University – Computer and Information Sciences")
    d.add_paragraph("Dear Editor,")
    body = [
        f"We submit the manuscript “{TITLE_V2}” for consideration as a research article.",
        "Many teams attach conformal prediction sets to their models to state how sure a prediction is, and they "
        "check the promised coverage on a random test split. We show that this check can be badly wrong when the "
        "data contain repeated rows, which is common in software-defect, network-intrusion and phishing data. "
        f"On the NASA mc1 defect data, for example, the usual check reports {n['cs_mc1_reported']}% coverage while "
        f"new modules are covered only {n['cs_mc1_new']}% of the time, often with an empty prediction set.",
        "The paper gives practitioners a short routine. A check that needs only the data predicts how large the "
        "problem can be before any model is trained; a fix (separate calibration of seen and new cases) restores "
        f"the promised coverage for new cases (on mc1, {n['cs_mc1_fixed']}%); and a reporting rule makes the problem "
        f"visible. We evaluate the routine on three applications, on {n['n_total']} standard benchmark data sets, "
        "in controlled experiments, and on a fresh set of data sets analysed under a plan fixed in advance. "
        + n["cover_prospective"],
        "We report negative results openly: on most benchmark data the problem is small, and the check correctly "
        "says so. All code and results are openly available, and one command reproduces every number.",
        "[AUTHOR TO CONFIRM: This manuscript is original, has not been published, and is not under consideration "
        "elsewhere. All authors have approved the submission. We declare no competing interests.]",
        "The manuscript was prepared with substantial assistance from an AI system, as described in its "
        "“Use of AI tools” statement.",
        "Sincerely,\n[CORRESPONDING AUTHOR NAME, AFFILIATION, E-MAIL]",
    ]
    for b in body:
        d.add_paragraph(b)
    d.save(os.path.join(HERE, "Revised_Cover_Letter.docx"))


def highlights_docx():
    d = _doc()
    d.add_heading("Highlights", level=1)
    for line in open(os.path.join(HERE, "highlights.txt"), encoding="utf8").read().splitlines():
        if line.strip():
            d.add_paragraph(line.strip(), style="List Bullet")
    d.save(os.path.join(HERE, "highlights_PatternRecognition.docx"))


if __name__ == "__main__":
    if sys.argv[1:] == ["v2"]:
        from numbers_v2 import compute_numbers_v2
        title_page_v2()
        cover_letter_v2(compute_numbers_v2(os.path.join(os.path.dirname(HERE), "results")))
        print("ok")
        sys.exit()
    nums = compute_numbers(os.path.join(os.path.dirname(HERE), "results"))
    title_page()
    cover_letter(nums)
    cover_letter_pr(nums)
    highlights_docx()
    print("ok")
